import os
import uuid
import shutil
from io import BytesIO
from typing import Optional

from fastapi import (
    APIRouter,
    UploadFile,
    File,
    HTTPException,
    Depends,
    Form,
)
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.config import UPLOAD_DIR
from app.services.player_client import get_player_data
from app.database import get_db
from app.models import Player
from app.schemas.player import PlayerResponse


router = APIRouter(prefix="/api", tags=["Players"])


# =========================================================
# VIDEO SETTINGS
# =========================================================

ALLOWED_EXTENSIONS = {
    ".mp4",
    ".avi",
    ".mov",
}

ALLOWED_MIME_TYPES = {
    "video/mp4",
    "video/x-msvideo",
    "video/quicktime",
}


# =========================================================
# IMAGE SETTINGS
# =========================================================

ALLOWED_IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".webp",
}

ALLOWED_IMAGE_MIME_TYPES = {
    "image/jpeg",
    "image/png",
    "image/gif",
    "image/webp",
}

PLAYER_PHOTOS_DIR = os.path.join(
    UPLOAD_DIR,
    "player_photos",
)


# =========================================================
# EXCEL SETTINGS
# =========================================================

ALLOWED_EXCEL_EXTENSIONS = {
    ".xlsx",
}

ALLOWED_EXCEL_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    # Sometimes browsers/clients send this:
    "application/octet-stream",
}


EXCEL_COLUMNS = {
    "name",
    "team",
    "position",
    "photo",
    "kicks",
    "handballs",
    "marks",
    "tackles",
    "goals",
    "efficiency",
    "age",
    "height",
    "weight",
    "jersey_number",
    "inside50s",
    "disposals",
    "team_logo",
    "notes",
}


REQUIRED_EXCEL_COLUMNS = {
    "name",
    "team",
    "position",
}


INTEGER_EXCEL_COLUMNS = {
    "kicks",
    "handballs",
    "marks",
    "tackles",
    "goals",
    "efficiency",
    "age",
    "jersey_number",
    "inside50s",
    "disposals",
}


# =========================================================
# HELPER FUNCTIONS
# =========================================================


def clean_string(value):
    """
    Convert Excel values into clean strings.
    Empty values become None.
    """

    if value is None:
        return None

    value = str(value).strip()

    if not value:
        return None

    return value


def excel_to_int(
    value,
    column_name: str,
    row_number: int,
    default: int = 0,
):
    """
    Safely convert an Excel value to integer.

    Raises ValueError if an invalid value is provided
    instead of silently converting invalid data to 0.
    """

    if value is None or value == "":
        return default

    # Excel may return numbers like 10.0
    if isinstance(value, float):
        if value.is_integer():
            return int(value)

        raise ValueError(
            f"Column '{column_name}' must be a whole number"
        )

    try:
        return int(value)

    except (ValueError, TypeError):
        raise ValueError(
            f"Column '{column_name}' must be a valid integer"
        )


# =========================================================
# PLAYER TRACKING FROM VIDEO
# =========================================================


@router.post("/players")
async def run_player_tracking(
    file: UploadFile = File(...)
):
    ext = os.path.splitext(
        file.filename or ""
    )[1].lower()

    if (
        ext not in ALLOWED_EXTENSIONS
        or file.content_type not in ALLOWED_MIME_TYPES
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid video format. "
                "Accepted: .mp4, .avi, .mov"
            ),
        )

    os.makedirs(
        UPLOAD_DIR,
        exist_ok=True,
    )

    tmp_path = os.path.join(
        UPLOAD_DIR,
        f"tmp_{uuid.uuid4()}{ext}",
    )

    try:
        with open(tmp_path, "wb") as f:
            shutil.copyfileobj(
                file.file,
                f,
            )

        data = await get_player_data(
            tmp_path
        )

        return {
            "status": "success",
            "data": data,
        }

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=str(e),
        )

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


# =========================================================
# CREATE SINGLE PLAYER
# =========================================================


@router.post(
    "/player",
    response_model=dict,
)
async def create_player(
    name: str = Form(...),
    team: str = Form(...),
    position: str = Form(...),
    photo: Optional[UploadFile] = File(None),
    kicks: int = Form(0),
    handballs: int = Form(0),
    marks: int = Form(0),
    tackles: int = Form(0),
    goals: int = Form(0),
    efficiency: int = Form(75),
    age: int = Form(0),
    height: Optional[str] = Form(None),
    weight: Optional[str] = Form(None),
    jerseyNumber: int = Form(0),
    inside50s: int = Form(0),
    disposals: int = Form(0),
    teamLogo: Optional[str] = Form(None),
    notes: Optional[str] = Form(None),
    db: Session = Depends(get_db),
):
    photo_path = None

    # -----------------------------------------------------
    # Handle player photo
    # -----------------------------------------------------

    if photo:
        ext = os.path.splitext(
            photo.filename or ""
        )[1].lower()

        if (
            ext not in ALLOWED_IMAGE_EXTENSIONS
            or photo.content_type
            not in ALLOWED_IMAGE_MIME_TYPES
        ):
            raise HTTPException(
                status_code=400,
                detail=(
                    "Invalid image format. "
                    "Accepted: .jpg, .jpeg, "
                    ".png, .gif, .webp"
                ),
            )

        os.makedirs(
            PLAYER_PHOTOS_DIR,
            exist_ok=True,
        )

        safe_name = name.replace(
            " ",
            "_",
        )

        unique_filename = (
            f"{uuid.uuid4()}_"
            f"{safe_name}{ext}"
        )

        photo_full_path = os.path.join(
            PLAYER_PHOTOS_DIR,
            unique_filename,
        )

        try:
            with open(
                photo_full_path,
                "wb",
            ) as buffer:
                shutil.copyfileobj(
                    photo.file,
                    buffer,
                )

            photo_path = os.path.join(
                "player_photos",
                unique_filename,
            )

        except Exception as e:
            raise HTTPException(
                status_code=500,
                detail=(
                    "Failed to save image: "
                    f"{str(e)}"
                ),
            )

    # -----------------------------------------------------
    # Save player to database
    # -----------------------------------------------------

    try:
        db_player = Player(
            name=name,
            team=team,
            position=position,
            photo=photo_path,
            kicks=kicks,
            handballs=handballs,
            marks=marks,
            tackles=tackles,
            goals=goals,
            efficiency=efficiency,
            age=age,
            height=height,
            weight=weight,
            jersey_number=jerseyNumber,
            inside50s=inside50s,
            disposals=disposals,
            team_logo=teamLogo,
            notes=notes,
        )

        db.add(db_player)

        db.commit()

        db.refresh(db_player)

        return {
            "message": (
                "Player created successfully"
            ),
            "player": {
                "id": db_player.id,
                "name": db_player.name,
                "team": db_player.team,
                "position": db_player.position,
                "photo": db_player.photo,
                "kicks": db_player.kicks,
                "handballs": db_player.handballs,
                "marks": db_player.marks,
                "tackles": db_player.tackles,
                "goals": db_player.goals,
                "efficiency": db_player.efficiency,
                "age": db_player.age,
                "height": db_player.height,
                "weight": db_player.weight,
                "jersey_number": (
                    db_player.jersey_number
                ),
                "inside50s": db_player.inside50s,
                "disposals": db_player.disposals,
                "team_logo": db_player.team_logo,
                "notes": db_player.notes,
            },
        }

    except Exception as e:
        db.rollback()

        # Delete saved photo if DB save failed
        if photo_path:
            full_path = os.path.join(
                UPLOAD_DIR,
                photo_path,
            )

            if os.path.exists(full_path):
                os.remove(full_path)

        raise HTTPException(
            status_code=500,
            detail=(
                "Failed to create player: "
                f"{str(e)}"
            ),
        )


# =========================================================
# IMPORT PLAYERS FROM EXCEL
# =========================================================


@router.post("/players/upload-excel")
async def upload_players_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """
    Upload an Excel file containing players.

    Required Excel columns:
        name
        team
        position

    Optional columns:
        photo
        kicks
        handballs
        marks
        tackles
        goals
        efficiency
        age
        height
        weight
        jersey_number
        inside50s
        disposals
        team_logo
        notes
    """

    # -----------------------------------------------------
    # Validate Excel extension
    # -----------------------------------------------------

    filename = file.filename or ""

    ext = os.path.splitext(
        filename
    )[1].lower()

    if ext not in ALLOWED_EXCEL_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid Excel file. "
                "Please upload a .xlsx file."
            ),
        )

    # -----------------------------------------------------
    # Read Excel file
    # -----------------------------------------------------

    try:
        contents = await file.read()

        if not contents:
            raise HTTPException(
                status_code=400,
                detail="The uploaded Excel file is empty.",
            )

        workbook = load_workbook(
            filename=BytesIO(contents),
            data_only=True,
            read_only=True,
        )

    except HTTPException:
        raise

    except Exception:
        raise HTTPException(
            status_code=400,
            detail=(
                "Could not read the Excel file. "
                "Please make sure it is a valid .xlsx file."
            ),
        )

    try:
        sheet = workbook.active

        # -------------------------------------------------
        # Validate headers
        # -------------------------------------------------

        first_row = next(
            sheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            ),
            None,
        )

        if not first_row:
            raise HTTPException(
                status_code=400,
                detail="Excel file has no header row.",
            )

        headers = []

        for header in first_row:
            if header is None:
                headers.append(None)
            else:
                headers.append(
                    str(header)
                    .strip()
                    .lower()
                )

        # Detect duplicate headers
        non_empty_headers = [
            header
            for header in headers
            if header
        ]

        duplicate_headers = {
            header
            for header in non_empty_headers
            if non_empty_headers.count(header) > 1
        }

        if duplicate_headers:
            raise HTTPException(
                status_code=400,
                detail={
                    "message": (
                        "Excel contains duplicate columns."
                    ),
                    "columns": sorted(
                        duplicate_headers
                    ),
                },
            )

        missing_columns = (
            REQUIRED_EXCEL_COLUMNS
            - set(non_empty_headers)
        )

        if missing_columns:
            raise HTTPException(
                status_code=400,
                detail={
                    "message": (
                        "Missing required Excel columns."
                    ),
                    "missing_columns": sorted(
                        missing_columns
                    ),
                },
            )

        unknown_columns = (
            set(non_empty_headers)
            - EXCEL_COLUMNS
        )

        if unknown_columns:
            raise HTTPException(
                status_code=400,
                detail={
                    "message": (
                        "Excel contains unsupported columns."
                    ),
                    "unsupported_columns": sorted(
                        unknown_columns
                    ),
                    "allowed_columns": sorted(
                        EXCEL_COLUMNS
                    ),
                },
            )

        # -------------------------------------------------
        # Validate ALL rows before inserting anything
        # -------------------------------------------------

        validated_players = []
        validation_errors = []

        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            # Skip completely empty rows
            if all(
                value is None
                or (
                    isinstance(value, str)
                    and not value.strip()
                )
                for value in row
            ):
                continue

            row_data = {}

            for index, header in enumerate(headers):
                if not header:
                    continue

                value = (
                    row[index]
                    if index < len(row)
                    else None
                )

                row_data[header] = value

            row_errors = []

            # ---------------------------------------------
            # Required values
            # ---------------------------------------------

            name = clean_string(
                row_data.get("name")
            )

            team = clean_string(
                row_data.get("team")
            )

            position = clean_string(
                row_data.get("position")
            )

            if not name:
                row_errors.append(
                    "'name' is required"
                )

            if not team:
                row_errors.append(
                    "'team' is required"
                )

            if not position:
                row_errors.append(
                    "'position' is required"
                )

            # ---------------------------------------------
            # Integer validation
            # ---------------------------------------------

            integer_values = {}

            defaults = {
                "kicks": 0,
                "handballs": 0,
                "marks": 0,
                "tackles": 0,
                "goals": 0,
                "efficiency": 75,
                "age": 0,
                "jersey_number": 0,
                "inside50s": 0,
                "disposals": 0,
            }

            for column in INTEGER_EXCEL_COLUMNS:
                try:
                    integer_values[column] = (
                        excel_to_int(
                            row_data.get(column),
                            column,
                            row_number,
                            defaults[column],
                        )
                    )

                except ValueError as e:
                    row_errors.append(
                        str(e)
                    )

            # ---------------------------------------------
            # Optional business validation
            # ---------------------------------------------

            if (
                "efficiency" in integer_values
                and not (
                    0
                    <= integer_values["efficiency"]
                    <= 100
                )
            ):
                row_errors.append(
                    "'efficiency' must be "
                    "between 0 and 100"
                )

            if (
                "age" in integer_values
                and integer_values["age"] < 0
            ):
                row_errors.append(
                    "'age' cannot be negative"
                )

            if (
                "jersey_number" in integer_values
                and integer_values["jersey_number"] < 0
            ):
                row_errors.append(
                    "'jersey_number' cannot be negative"
                )

            stats_to_check = [
                "kicks",
                "handballs",
                "marks",
                "tackles",
                "goals",
                "inside50s",
                "disposals",
            ]

            for stat in stats_to_check:
                if (
                    stat in integer_values
                    and integer_values[stat] < 0
                ):
                    row_errors.append(
                        f"'{stat}' cannot be negative"
                    )

            # ---------------------------------------------
            # If row contains errors
            # ---------------------------------------------

            if row_errors:
                validation_errors.append(
                    {
                        "row": row_number,
                        "errors": row_errors,
                    }
                )

                continue

            # ---------------------------------------------
            # Build validated player data
            # ---------------------------------------------

            validated_players.append(
                {
                    "excel_row": row_number,
                    "name": name,
                    "team": team,
                    "position": position,
                    "photo": clean_string(
                        row_data.get("photo")
                    ),
                    "kicks": integer_values["kicks"],
                    "handballs": integer_values[
                        "handballs"
                    ],
                    "marks": integer_values["marks"],
                    "tackles": integer_values[
                        "tackles"
                    ],
                    "goals": integer_values["goals"],
                    "efficiency": integer_values[
                        "efficiency"
                    ],
                    "age": integer_values["age"],
                    "height": clean_string(
                        row_data.get("height")
                    ),
                    "weight": clean_string(
                        row_data.get("weight")
                    ),
                    "jersey_number": integer_values[
                        "jersey_number"
                    ],
                    "inside50s": integer_values[
                        "inside50s"
                    ],
                    "disposals": integer_values[
                        "disposals"
                    ],
                    "team_logo": clean_string(
                        row_data.get("team_logo")
                    ),
                    "notes": clean_string(
                        row_data.get("notes")
                    ),
                }
            )

        # -------------------------------------------------
        # No players found
        # -------------------------------------------------

        if (
            not validated_players
            and not validation_errors
        ):
            raise HTTPException(
                status_code=400,
                detail=(
                    "The Excel file does not contain "
                    "any player rows."
                ),
            )

        # -------------------------------------------------
        # Reject entire upload if any row is invalid
        # -------------------------------------------------

        if validation_errors:
            raise HTTPException(
                status_code=422,
                detail={
                    "message": (
                        "Excel validation failed. "
                        "No players were imported."
                    ),
                    "total_errors": len(
                        validation_errors
                    ),
                    "errors": validation_errors,
                },
            )

        # -------------------------------------------------
        # Insert players
        # -------------------------------------------------

        created_players = []

        try:
            for player_data in validated_players:
                excel_row = player_data.pop(
                    "excel_row"
                )

                db_player = Player(
                    **player_data
                )

                db.add(db_player)

                # Gets database-generated ID
                # without committing yet
                db.flush()

                created_players.append(
                    {
                        "excel_row": excel_row,
                        "id": db_player.id,
                        "name": db_player.name,
                        "team": db_player.team,
                        "position": db_player.position,
                    }
                )

            # Commit all players together
            db.commit()

        except Exception as e:
            db.rollback()

            raise HTTPException(
                status_code=500,
                detail=(
                    "Failed to save players "
                    f"to database: {str(e)}"
                ),
            )

        return {
            "message": (
                "Players imported successfully"
            ),
            "players_imported": len(
                created_players
            ),
            "players": created_players,
        }

    finally:
        workbook.close()


# =========================================================
# GET ALL PLAYERS
# =========================================================


@router.get(
    "/players",
    response_model=list[PlayerResponse],
)
def get_players(
    db: Session = Depends(get_db)
):
    players = (
        db.query(Player)
        .all()
    )

    return players


# =========================================================
# GET SINGLE PLAYER
# =========================================================


@router.get(
    "/player/{player_id}",
    response_model=PlayerResponse,
)
def get_player(
    player_id: int,
    db: Session = Depends(get_db),
):
    player = (
        db.query(Player)
        .filter(
            Player.id == player_id
        )
        .first()
    )

    if player is None:
        raise HTTPException(
            status_code=404,
            detail="Player not found",
        )

    return player


# =========================================================
# DELETE PLAYER
# =========================================================


@router.delete(
    "/player/{player_id}"
)
def delete_player(
    player_id: int,
    db: Session = Depends(get_db),
):
    player = (
        db.query(Player)
        .filter(
            Player.id == player_id
        )
        .first()
    )

    if player is None:
        raise HTTPException(
            status_code=404,
            detail="Player not found",
        )

    # -----------------------------------------------------
    # Delete photo file if it exists
    # -----------------------------------------------------

    if player.photo:
        photo_full_path = os.path.join(
            UPLOAD_DIR,
            player.photo,
        )

        if os.path.exists(
            photo_full_path
        ):
            try:
                os.remove(
                    photo_full_path
                )

            except Exception as e:
                print(
                    "Warning: Could not delete "
                    f"photo file: {str(e)}"
                )

    # -----------------------------------------------------
    # Delete database record
    # -----------------------------------------------------

    try:
        db.delete(player)

        db.commit()

    except Exception as e:
        db.rollback()

        raise HTTPException(
            status_code=500,
            detail=(
                "Failed to delete player: "
                f"{str(e)}"
            ),
        )

    return {
        "message": (
            "Player deleted successfully"
        ),
        "player_id": player_id,
    }